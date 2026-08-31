import io
import json
from copy import copy
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field

from .auth import request_has_role
from .db import connect, now_iso
from .rules import BusinessError

router = APIRouter(prefix="/api/investments", tags=["数字化投入管理"])

PLAN_NODES = ("部门负责人审批", "财务审批", "分管领导审批")
PLAN_ROLES = {
    "部门负责人审批": ("department_head", "admin"),
    "财务审批": ("finance", "admin"),
    "分管领导审批": ("vp", "business_owner", "admin"),
}
ADJUST_ROLES = {
    "部门负责人审批": ("department_head", "admin"),
    "财务审批": ("finance", "admin"),
    "分管领导审批": ("vp", "business_owner", "admin"),
}


class CategoryPayload(BaseModel):
    category_name: str = Field(min_length=1, max_length=80)
    subcategory_name: str = Field(min_length=1, max_length=80)
    tags: list[str] = Field(default_factory=list, max_length=20)
    status: str = "启用"
    sort_order: int = Field(default=0, ge=0, le=9999)


class PlanPayload(BaseModel):
    plan_name: str = Field(min_length=1, max_length=160)
    plan_year: int = Field(ge=2020, le=2100)
    department: str = Field(min_length=1, max_length=100)
    description: str = Field(default="", max_length=3000)
    prior_year_budget: float = Field(default=0, ge=0)
    prior_year_actual: float = Field(default=0, ge=0)
    current_year_budget: float = Field(default=0, ge=0)
    current_year_actual: float = Field(default=0, ge=0)


class ItemPayload(BaseModel):
    item_name: str = Field(min_length=1, max_length=160)
    is_new: bool = False
    category_id: Optional[int] = None
    category_name: str = Field(default="", max_length=80)
    subcategory_name: str = Field(default="", max_length=80)
    custom_tags: list[str] = Field(default_factory=list, max_length=20)
    quantity: float = Field(default=1, gt=0)
    unit: str = Field(default="项", max_length=20)
    application_amount: float = Field(gt=0)
    approved_amount: Optional[float] = Field(default=None, ge=0)
    payer: str = Field(min_length=1, max_length=120)
    business_purpose: str = Field(min_length=1, max_length=1000)
    is_unplanned_reserve: bool = False
    project_id: Optional[int] = None
    contract_id: Optional[int] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    planned_payment_amount: float = Field(default=0, ge=0)


class ApprovalPayload(BaseModel):
    action: str = "通过"
    comment: str = Field(default="", max_length=1000)


class BatchApprovalPayload(ApprovalPayload):
    ids: list[int] = Field(min_length=1, max_length=100)


class AdjustmentPayload(BaseModel):
    plan_id: int
    item_id: int
    adjustment_type: str = Field(default="金额与范围调整", max_length=80)
    requested_amount: float = Field(ge=0)
    scope_after: str = Field(default="", max_length=2000)
    reason: str = Field(min_length=1, max_length=2000)


class BindingPayload(BaseModel):
    project_id: Optional[int] = None
    contract_id: Optional[int] = None
    planned_payment_amount: Optional[float] = Field(default=None, ge=0)


class PaymentPayload(BaseModel):
    item_id: int
    contract_id: Optional[int] = None
    payment_type: str = Field(pattern=r"^(合同付款|普通费用)$")
    payment_year: int = Field(ge=2020, le=2100)
    amount: float = Field(gt=0)
    payment_date: str
    document_no: str = Field(min_length=1, max_length=120)
    payer: str = Field(min_length=1, max_length=120)
    description: str = Field(default="", max_length=1000)


class WarningRulePayload(BaseModel):
    threshold_value: float = Field(ge=0, le=1000)
    days_value: int = Field(default=0, ge=0, le=3650)
    enabled: bool = True
    level: str = Field(pattern=r"^(提示|预警|严重)$")


def _actor(request: Request) -> tuple[str, str]:
    user = getattr(request.state, "auth_user", None) or {}
    if user:
        return user.get("display_name") or user.get("username") or "系统用户", user.get("role_code") or "applicant"
    return request.headers.get("X-User", "测试用户"), request.headers.get("X-Role", "applicant")


def _require_role(request: Request, *roles: str):
    if not request_has_role(request, *roles):
        raise BusinessError(403, "INV-4030", "当前角色无权处理该投入业务")


def _iso_date(value: Optional[str], label: str) -> Optional[str]:
    if not value:
        return None
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise BusinessError(400, "INV-4001", f"{label}必须为 YYYY-MM-DD") from exc


def _next_no(conn, table: str, column: str, prefix: str) -> str:
    row = conn.execute(f"SELECT {column} no FROM {table} WHERE {column} LIKE ? ORDER BY id DESC LIMIT 1", (f"{prefix}%",)).fetchone()
    seq = 1
    if row:
        try:
            seq = int(str(row["no"]).rsplit("-", 1)[1]) + 1
        except (ValueError, IndexError):
            pass
    return f"{prefix}{seq:04d}"


def _json_list(raw) -> list:
    try:
        value = json.loads(raw or "[]")
        return value if isinstance(value, list) else []
    except Exception:
        return []


def init_investment_db():
    with connect() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS investment_categories (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                category_name TEXT NOT NULL,
                subcategory_name TEXT NOT NULL,
                tags_json TEXT NOT NULL DEFAULT '[]',
                status TEXT NOT NULL DEFAULT '启用',
                sort_order INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(category_name,subcategory_name)
            );
            CREATE TABLE IF NOT EXISTS investment_plans (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                plan_no TEXT UNIQUE NOT NULL,
                plan_name TEXT NOT NULL,
                plan_year INTEGER NOT NULL,
                department TEXT NOT NULL,
                applicant TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT '草稿',
                current_node TEXT NOT NULL DEFAULT '草稿',
                prior_year_budget REAL NOT NULL DEFAULT 0,
                prior_year_actual REAL NOT NULL DEFAULT 0,
                current_year_budget REAL NOT NULL DEFAULT 0,
                current_year_actual REAL NOT NULL DEFAULT 0,
                application_total REAL NOT NULL DEFAULT 0,
                approved_total REAL NOT NULL DEFAULT 0,
                description TEXT DEFAULT '',
                submitted_at TEXT,
                approved_at TEXT,
                finance_confirmed_at TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS investment_items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                plan_id INTEGER NOT NULL,
                item_no TEXT UNIQUE NOT NULL,
                item_name TEXT NOT NULL,
                is_new INTEGER NOT NULL DEFAULT 0,
                category_id INTEGER,
                category_name TEXT NOT NULL,
                subcategory_name TEXT NOT NULL,
                custom_tags_json TEXT NOT NULL DEFAULT '[]',
                quantity REAL NOT NULL DEFAULT 1,
                unit TEXT NOT NULL DEFAULT '项',
                application_amount REAL NOT NULL,
                approved_amount REAL NOT NULL DEFAULT 0,
                payer TEXT NOT NULL,
                business_purpose TEXT NOT NULL,
                is_unplanned_reserve INTEGER NOT NULL DEFAULT 0,
                project_id INTEGER,
                contract_id INTEGER,
                start_date TEXT,
                end_date TEXT,
                planned_payment_amount REAL NOT NULL DEFAULT 0,
                paid_amount REAL NOT NULL DEFAULT 0,
                written_off_amount REAL NOT NULL DEFAULT 0,
                progress REAL NOT NULL DEFAULT 0,
                status TEXT NOT NULL DEFAULT '规划中',
                baseline_version INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(plan_id) REFERENCES investment_plans(id) ON DELETE CASCADE,
                FOREIGN KEY(category_id) REFERENCES investment_categories(id),
                FOREIGN KEY(project_id) REFERENCES projects(id),
                FOREIGN KEY(contract_id) REFERENCES contracts(id)
            );
            CREATE TABLE IF NOT EXISTS investment_approvals (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                plan_id INTEGER NOT NULL,
                node TEXT NOT NULL,
                role TEXT NOT NULL,
                approver TEXT NOT NULL,
                action TEXT NOT NULL,
                comment TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                FOREIGN KEY(plan_id) REFERENCES investment_plans(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS investment_adjustments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                adjustment_no TEXT UNIQUE NOT NULL,
                plan_id INTEGER NOT NULL,
                item_id INTEGER NOT NULL,
                adjustment_type TEXT NOT NULL,
                original_amount REAL NOT NULL,
                requested_amount REAL NOT NULL,
                amount_delta REAL NOT NULL,
                scope_before TEXT DEFAULT '',
                scope_after TEXT DEFAULT '',
                reason TEXT NOT NULL,
                applicant TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT '草稿',
                current_node TEXT NOT NULL DEFAULT '草稿',
                submitted_at TEXT,
                approved_at TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(plan_id) REFERENCES investment_plans(id) ON DELETE CASCADE,
                FOREIGN KEY(item_id) REFERENCES investment_items(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS investment_adjustment_approvals (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                adjustment_id INTEGER NOT NULL,
                node TEXT NOT NULL,
                role TEXT NOT NULL,
                approver TEXT NOT NULL,
                action TEXT NOT NULL,
                comment TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                FOREIGN KEY(adjustment_id) REFERENCES investment_adjustments(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS investment_payments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                payment_no TEXT UNIQUE NOT NULL,
                item_id INTEGER NOT NULL,
                contract_id INTEGER,
                payment_type TEXT NOT NULL,
                investment_year INTEGER NOT NULL,
                payment_year INTEGER NOT NULL,
                amount REAL NOT NULL,
                payment_date TEXT NOT NULL,
                document_no TEXT UNIQUE NOT NULL,
                payer TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT '已核销',
                writeoff_amount REAL NOT NULL,
                description TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(item_id) REFERENCES investment_items(id) ON DELETE CASCADE,
                FOREIGN KEY(contract_id) REFERENCES contracts(id)
            );
            CREATE TABLE IF NOT EXISTS investment_warning_rules (
                code TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                threshold_value REAL NOT NULL DEFAULT 0,
                days_value INTEGER NOT NULL DEFAULT 0,
                enabled INTEGER NOT NULL DEFAULT 1,
                level TEXT NOT NULL DEFAULT '预警',
                description TEXT DEFAULT '',
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS investment_warnings (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_key TEXT UNIQUE NOT NULL,
                plan_id INTEGER,
                item_id INTEGER,
                rule_code TEXT NOT NULL,
                level TEXT NOT NULL,
                title TEXT NOT NULL,
                content TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT '待处理',
                triggered_at TEXT NOT NULL,
                resolved_at TEXT,
                FOREIGN KEY(plan_id) REFERENCES investment_plans(id) ON DELETE CASCADE,
                FOREIGN KEY(item_id) REFERENCES investment_items(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_investment_plan_status ON investment_plans(status,current_node);
            CREATE INDEX IF NOT EXISTS idx_investment_item_plan ON investment_items(plan_id,status);
            CREATE INDEX IF NOT EXISTS idx_investment_payment_item ON investment_payments(item_id,payment_date);
            CREATE INDEX IF NOT EXISTS idx_investment_warning_status ON investment_warnings(status,triggered_at);
            """
        )
        ts = now_iso()
        defaults = [
            ("软件与平台", "应用系统建设", ["建设", "软件"], 10),
            ("软件与平台", "软件订阅与许可", ["SaaS", "订阅"], 20),
            ("基础设施", "云资源与算力", ["云服务", "算力"], 30),
            ("基础设施", "硬件设备", ["设备", "采购"], 40),
            ("数据与智能", "数据治理", ["数据", "治理"], 50),
            ("数据与智能", "AI与模型服务", ["AI", "模型"], 60),
            ("运维与安全", "运维服务", ["运维", "服务"], 70),
            ("运维与安全", "网络安全", ["安全", "合规"], 80),
        ]
        for category, subcategory, tags, sort_order in defaults:
            conn.execute(
                """INSERT INTO investment_categories(category_name,subcategory_name,tags_json,status,sort_order,created_at,updated_at)
                   VALUES(?,?,?,?,?,?,?) ON CONFLICT(category_name,subcategory_name) DO NOTHING""",
                (category, subcategory, json.dumps(tags, ensure_ascii=False), "启用", sort_order, ts, ts),
            )
        rules = [
            ("near_budget", "预算执行接近上限", 90, 0, "预警", "已核销金额达到审核后金额阈值"),
            ("over_budget", "超预算", 100, 0, "严重", "已核销金额超过审核后金额"),
            ("long_unexecuted", "长期未执行", 0, 90, "预警", "投入项生效后长期无付款核销"),
            ("near_expiry", "投入项临近到期", 0, 30, "提示", "距投入项结束日期小于阈值"),
            ("payment_deviation", "付款进度偏离计划", 20, 0, "预警", "实际付款进度与计划进度偏差超过阈值"),
        ]
        for code, name, threshold, days, level, description in rules:
            conn.execute(
                """INSERT INTO investment_warning_rules(code,name,threshold_value,days_value,enabled,level,description,updated_at)
                   VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(code) DO NOTHING""",
                (code, name, threshold, days, 1, level, description, ts),
            )
        # Editable POC ledger: these are ordinary database rows (not frontend
        # constants) and are inserted only into a brand-new investment schema.
        if conn.execute("SELECT COUNT(*) c FROM investment_plans").fetchone()["c"] == 0:
            current_year = date.today().year
            next_year = current_year + 1
            conn.execute(
                """INSERT INTO investment_plans(plan_no,plan_name,plan_year,department,applicant,status,current_node,
                   prior_year_budget,prior_year_actual,current_year_budget,current_year_actual,application_total,approved_total,
                   description,submitted_at,approved_at,finance_confirmed_at,created_at,updated_at)
                   VALUES(?,?,?,?,?,'已生效','已生效',?,?,?,?,?,?,?,?,?,?,?,?)""",
                (f"INV-{current_year}-0001", f"{current_year}年数字化核心能力投入计划", current_year, "数字化管理部", "李莉",
                 3200000, 2760000, 4200000, 1920000, 1500000, 1460000, "围绕智能化平台、数据治理与安全合规持续建设。", ts, ts, ts, ts, ts),
            )
            effective_id = conn.execute("SELECT id FROM investment_plans WHERE plan_no=?", (f"INV-{current_year}-0001",)).fetchone()["id"]
            project = conn.execute("SELECT id FROM projects ORDER BY id LIMIT 1").fetchone()
            contract = conn.execute("SELECT id FROM contracts ORDER BY id LIMIT 1").fetchone()
            cat_ai = conn.execute("SELECT id FROM investment_categories WHERE subcategory_name='AI与模型服务'").fetchone()
            cat_data = conn.execute("SELECT id FROM investment_categories WHERE subcategory_name='数据治理'").fetchone()
            conn.executemany(
                """INSERT INTO investment_items(plan_id,item_no,item_name,is_new,category_id,category_name,subcategory_name,
                   custom_tags_json,quantity,unit,application_amount,approved_amount,payer,business_purpose,is_unplanned_reserve,
                   project_id,contract_id,start_date,end_date,planned_payment_amount,paid_amount,written_off_amount,progress,status,
                   baseline_version,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                [
                    (effective_id, f"INVITEM-{current_year}-0001", "企业智能体与知识库服务", 1, cat_ai["id"], "数据与智能", "AI与模型服务", json.dumps(["重点建设", "AI"], ensure_ascii=False), 1, "年", 900000, 880000, "数字化管理部", "建设企业级智能问答与业务智能体能力。", 0, project["id"] if project else None, contract["id"] if contract else None, f"{current_year}-03-01", f"{current_year}-12-15", 700000, 420000, 420000, 47.73, "执行中", 1, ts, ts),
                    (effective_id, f"INVITEM-{current_year}-0002", "数据资产治理与质量提升", 0, cat_data["id"], "数据与智能", "数据治理", json.dumps(["数据资产"], ensure_ascii=False), 1, "项", 600000, 580000, "科技管理部", "完善数据标准、血缘与质量监控。", 0, None, None, f"{current_year}-05-01", f"{current_year}-11-30", 420000, 180000, 180000, 31.03, "执行中", 1, ts, ts),
                ],
            )
            first_item = conn.execute("SELECT id FROM investment_items WHERE item_no=?", (f"INVITEM-{current_year}-0001",)).fetchone()["id"]
            conn.execute(
                """INSERT INTO investment_payments(payment_no,item_id,contract_id,payment_type,investment_year,payment_year,
                   amount,payment_date,document_no,payer,status,writeoff_amount,description,created_at,updated_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,'已核销',?,?,?,?)""",
                (f"PAY-{current_year}-0001", first_item, contract["id"] if contract else None, "合同付款" if contract else "普通费用", current_year, current_year,
                 420000, f"{current_year}-06-30", f"POC-PAY-{current_year}-001", "数字化管理部", 420000, "年度服务第一阶段核销", ts, ts),
            )
            conn.execute(
                """INSERT INTO investment_plans(plan_no,plan_name,plan_year,department,applicant,status,current_node,
                   prior_year_budget,prior_year_actual,current_year_budget,current_year_actual,description,created_at,updated_at)
                   VALUES(?,?,?,?,?,'草稿','草稿',?,?,?,?,?,?,?)""",
                (f"INV-{next_year}-0001", f"{next_year}年数字化投入规划", next_year, "数字化管理部", "李莉", 4200000, 1920000, 5000000, 0, "参考上年预算及当年执行情况编制的下一年度投入草案。", ts, ts),
            )


def _sync_plan_totals(conn, plan_id: int):
    total = conn.execute(
        "SELECT COALESCE(SUM(application_amount),0) app,COALESCE(SUM(approved_amount),0) approved FROM investment_items WHERE plan_id=?",
        (plan_id,),
    ).fetchone()
    conn.execute(
        "UPDATE investment_plans SET application_total=?,approved_total=?,updated_at=? WHERE id=?",
        (total["app"], total["approved"], now_iso(), plan_id),
    )


def _item_dict(row) -> dict:
    d = dict(row)
    d["custom_tags"] = _json_list(d.pop("custom_tags_json", "[]"))
    d["is_new"] = bool(d.get("is_new"))
    d["is_unplanned_reserve"] = bool(d.get("is_unplanned_reserve"))
    d["remaining_amount"] = max(0, float(d.get("approved_amount") or 0) - float(d.get("written_off_amount") or 0))
    d["execution_rate"] = round(float(d.get("written_off_amount") or 0) / float(d.get("approved_amount") or 1) * 100, 2) if d.get("approved_amount") else 0
    return d


def _plan_dict(conn, row, detail=False) -> dict:
    d = dict(row)
    if detail:
        d["items"] = [_item_dict(r) for r in conn.execute(
            """SELECT i.*,p.project_no,p.name project_name,c.contract_no,c.name contract_name
                 FROM investment_items i LEFT JOIN projects p ON p.id=i.project_id
                 LEFT JOIN contracts c ON c.id=i.contract_id WHERE i.plan_id=? ORDER BY i.id""", (d["id"],)
        )]
        d["approvals"] = [dict(r) for r in conn.execute("SELECT * FROM investment_approvals WHERE plan_id=? ORDER BY id", (d["id"],))]
    return d


def _insert_warning(conn, event_key, plan_id, item_id, rule_code, level, title, content, active_keys):
    active_keys.add(event_key)
    existing = conn.execute("SELECT id,status FROM investment_warnings WHERE event_key=?", (event_key,)).fetchone()
    if existing:
        conn.execute(
            "UPDATE investment_warnings SET level=?,title=?,content=?,status='待处理',resolved_at=NULL WHERE id=?",
            (level, title, content, existing["id"]),
        )
    else:
        conn.execute(
            """INSERT INTO investment_warnings(event_key,plan_id,item_id,rule_code,level,title,content,status,triggered_at)
               VALUES(?,?,?,?,?,?,?,'待处理',?)""",
            (event_key, plan_id, item_id, rule_code, level, title, content, now_iso()),
        )


def refresh_investment_warnings(conn):
    rules = {r["code"]: dict(r) for r in conn.execute("SELECT * FROM investment_warning_rules WHERE enabled=1")}
    active_keys = set()
    today = date.today()
    rows = conn.execute(
        """SELECT i.*,p.plan_no,p.plan_year,p.status plan_status FROM investment_items i
           JOIN investment_plans p ON p.id=i.plan_id WHERE p.status='已生效' AND i.status!='已取消'"""
    ).fetchall()
    for row in rows:
        amount = float(row["approved_amount"] or 0)
        used = float(row["written_off_amount"] or 0)
        rate = used / amount * 100 if amount else 0
        if "over_budget" in rules and rate > rules["over_budget"]["threshold_value"]:
            r = rules["over_budget"]
            _insert_warning(conn, f"over_budget:{row['id']}", row["plan_id"], row["id"], r["code"], r["level"], "投入项超预算", f"{row['item_no']} {row['item_name']}执行率{rate:.1f}%，已超过审核后金额。", active_keys)
        elif "near_budget" in rules and rate >= rules["near_budget"]["threshold_value"]:
            r = rules["near_budget"]
            _insert_warning(conn, f"near_budget:{row['id']}", row["plan_id"], row["id"], r["code"], r["level"], "预算执行接近上限", f"{row['item_no']} {row['item_name']}执行率已达{rate:.1f}%。", active_keys)
        if "near_expiry" in rules and row["end_date"]:
            try:
                remain_days = (date.fromisoformat(row["end_date"]) - today).days
                r = rules["near_expiry"]
                if 0 <= remain_days <= r["days_value"] and row["status"] != "已完成":
                    _insert_warning(conn, f"near_expiry:{row['id']}", row["plan_id"], row["id"], r["code"], r["level"], "投入项临近到期", f"{row['item_no']} {row['item_name']}距到期仅{remain_days}天。", active_keys)
            except ValueError:
                pass
        if "long_unexecuted" in rules and not used:
            try:
                created_days = (today - date.fromisoformat(str(row["created_at"])[:10])).days
                r = rules["long_unexecuted"]
                if created_days >= r["days_value"]:
                    _insert_warning(conn, f"long_unexecuted:{row['id']}", row["plan_id"], row["id"], r["code"], r["level"], "投入项长期未执行", f"{row['item_no']} {row['item_name']}生效后{created_days}天仍无核销。", active_keys)
            except ValueError:
                pass
        if "payment_deviation" in rules and row["planned_payment_amount"]:
            planned = min(100.0, float(row["planned_payment_amount"] or 0) / max(amount, 1) * 100)
            deviation = abs(planned - rate)
            r = rules["payment_deviation"]
            if deviation > r["threshold_value"]:
                _insert_warning(conn, f"payment_deviation:{row['id']}", row["plan_id"], row["id"], r["code"], r["level"], "付款进度偏离计划", f"{row['item_no']}计划付款占比{planned:.1f}%，实际核销{rate:.1f}%，偏差{deviation:.1f}%。", active_keys)
    open_rows = conn.execute("SELECT id,event_key FROM investment_warnings WHERE status='待处理'").fetchall()
    for warning in open_rows:
        if warning["event_key"] not in active_keys:
            conn.execute("UPDATE investment_warnings SET status='已恢复',resolved_at=? WHERE id=?", (now_iso(), warning["id"]))


@router.get("/categories")
def list_categories():
    with connect() as conn:
        data = []
        for row in conn.execute("SELECT * FROM investment_categories ORDER BY sort_order,id"):
            d = dict(row); d["tags"] = _json_list(d.pop("tags_json", "[]")); data.append(d)
    return {"code": 0, "data": data}


@router.post("/categories")
def create_category(payload: CategoryPayload, request: Request):
    _require_role(request, "finance", "admin")
    with connect() as conn:
        ts = now_iso()
        try:
            cur = conn.execute(
                """INSERT INTO investment_categories(category_name,subcategory_name,tags_json,status,sort_order,created_at,updated_at)
                   VALUES(?,?,?,?,?,?,?)""",
                (payload.category_name, payload.subcategory_name, json.dumps(payload.tags, ensure_ascii=False), payload.status, payload.sort_order, ts, ts),
            )
        except Exception as exc:
            raise BusinessError(409, "INV-4090", "该投入类别和子类别已存在") from exc
    return {"code": 0, "message": "分类标签已创建", "data": {"id": cur.lastrowid}}


@router.put("/categories/{category_id}")
def update_category(category_id: int, payload: CategoryPayload, request: Request):
    _require_role(request, "finance", "admin")
    with connect() as conn:
        if not conn.execute("SELECT 1 FROM investment_categories WHERE id=?", (category_id,)).fetchone():
            raise BusinessError(404, "INV-4040", "分类不存在")
        conn.execute(
            """UPDATE investment_categories SET category_name=?,subcategory_name=?,tags_json=?,status=?,sort_order=?,updated_at=? WHERE id=?""",
            (payload.category_name, payload.subcategory_name, json.dumps(payload.tags, ensure_ascii=False), payload.status, payload.sort_order, now_iso(), category_id),
        )
    return {"code": 0, "message": "分类标签已更新"}


@router.get("/plans")
def list_plans(year: Optional[int] = None, status: str = "", keyword: str = ""):
    where, params = ["1=1"], []
    if year: where.append("plan_year=?"); params.append(year)
    if status: where.append("status=?"); params.append(status)
    if keyword: where.append("(plan_no LIKE ? OR plan_name LIKE ? OR department LIKE ?)"); params.extend([f"%{keyword}%"] * 3)
    with connect() as conn:
        rows = conn.execute(f"SELECT * FROM investment_plans WHERE {' AND '.join(where)} ORDER BY plan_year DESC,id DESC", params).fetchall()
        data = [_plan_dict(conn, row) for row in rows]
    return {"code": 0, "data": data}


@router.post("/plans")
def create_plan(payload: PlanPayload, request: Request):
    actor, _ = _actor(request)
    if payload.prior_year_actual > payload.prior_year_budget and payload.prior_year_budget:
        raise BusinessError(422, "INV-4220", "上年实际执行额不应大于上年预算，请确认历史数据")
    with connect() as conn:
        ts = now_iso(); no = _next_no(conn, "investment_plans", "plan_no", f"INV-{payload.plan_year}-")
        cur = conn.execute(
            """INSERT INTO investment_plans(plan_no,plan_name,plan_year,department,applicant,status,current_node,
               prior_year_budget,prior_year_actual,current_year_budget,current_year_actual,description,created_at,updated_at)
               VALUES(?,?,?,?,?,'草稿','草稿',?,?,?,?,?,?,?)""",
            (no, payload.plan_name, payload.plan_year, payload.department, actor, payload.prior_year_budget, payload.prior_year_actual,
             payload.current_year_budget, payload.current_year_actual, payload.description, ts, ts),
        )
    return {"code": 0, "message": "投入计划已创建", "data": {"id": cur.lastrowid, "plan_no": no}}


@router.get("/plans/{plan_id}")
def get_plan(plan_id: int):
    with connect() as conn:
        row = conn.execute("SELECT * FROM investment_plans WHERE id=?", (plan_id,)).fetchone()
        if not row: raise BusinessError(404, "INV-4040", "投入计划不存在")
        data = _plan_dict(conn, row, True)
    return {"code": 0, "data": data}


@router.put("/plans/{plan_id}")
def update_plan(plan_id: int, payload: PlanPayload):
    with connect() as conn:
        row = conn.execute("SELECT status FROM investment_plans WHERE id=?", (plan_id,)).fetchone()
        if not row: raise BusinessError(404, "INV-4040", "投入计划不存在")
        if row["status"] not in ("草稿", "已驳回"): raise BusinessError(409, "INV-4091", "只有草稿或已驳回计划可编辑")
        conn.execute(
            """UPDATE investment_plans SET plan_name=?,plan_year=?,department=?,prior_year_budget=?,prior_year_actual=?,
               current_year_budget=?,current_year_actual=?,description=?,updated_at=? WHERE id=?""",
            (payload.plan_name, payload.plan_year, payload.department, payload.prior_year_budget, payload.prior_year_actual,
             payload.current_year_budget, payload.current_year_actual, payload.description, now_iso(), plan_id),
        )
    return {"code": 0, "message": "投入计划已保存"}


@router.post("/plans/{plan_id}/items")
def add_item(plan_id: int, payload: ItemPayload):
    start, end = _iso_date(payload.start_date, "开始日期"), _iso_date(payload.end_date, "结束日期")
    if start and end and start > end: raise BusinessError(400, "INV-4001", "投入项开始日期不能晚于结束日期")
    with connect() as conn:
        plan = conn.execute("SELECT * FROM investment_plans WHERE id=?", (plan_id,)).fetchone()
        if not plan: raise BusinessError(404, "INV-4040", "投入计划不存在")
        if plan["status"] not in ("草稿", "已驳回"): raise BusinessError(409, "INV-4091", "审批中或已生效计划不能直接修改明细")
        category_name, subcategory_name = payload.category_name, payload.subcategory_name
        if payload.category_id:
            category = conn.execute("SELECT * FROM investment_categories WHERE id=? AND status='启用'", (payload.category_id,)).fetchone()
            if not category: raise BusinessError(400, "INV-4001", "投入分类不存在或已停用")
            category_name, subcategory_name = category["category_name"], category["subcategory_name"]
        if not category_name or not subcategory_name: raise BusinessError(400, "INV-4001", "请选择投入类别和子类别")
        ts = now_iso(); item_no = _next_no(conn, "investment_items", "item_no", f"INVITEM-{plan['plan_year']}-")
        approved = payload.application_amount if payload.approved_amount is None else payload.approved_amount
        cur = conn.execute(
            """INSERT INTO investment_items(plan_id,item_no,item_name,is_new,category_id,category_name,subcategory_name,
               custom_tags_json,quantity,unit,application_amount,approved_amount,payer,business_purpose,is_unplanned_reserve,
               project_id,contract_id,start_date,end_date,planned_payment_amount,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (plan_id, item_no, payload.item_name, int(payload.is_new), payload.category_id, category_name, subcategory_name,
             json.dumps(payload.custom_tags, ensure_ascii=False), payload.quantity, payload.unit, payload.application_amount, approved,
             payload.payer, payload.business_purpose, int(payload.is_unplanned_reserve), payload.project_id, payload.contract_id,
             start, end, payload.planned_payment_amount, ts, ts),
        )
        _sync_plan_totals(conn, plan_id)
    return {"code": 0, "message": "投入明细已新增", "data": {"id": cur.lastrowid, "item_no": item_no}}


@router.get("/items")
def list_all_items(year: Optional[int] = None, status: str = "", keyword: str = ""):
    """All submitted/detail ledger rows, including draft and reserve items."""
    where, params = ["1=1"], []
    if year: where.append("p.plan_year=?"); params.append(year)
    if status: where.append("p.status=?"); params.append(status)
    if keyword:
        where.append("(i.item_no LIKE ? OR i.item_name LIKE ? OR p.plan_no LIKE ? OR p.department LIKE ?)")
        params.extend([f"%{keyword}%"] * 4)
    with connect() as conn:
        rows = conn.execute(
            f"""SELECT i.*,p.plan_no,p.plan_name,p.plan_year,p.department,p.status plan_status,p.current_node,
                pr.project_no,pr.name project_name,c.contract_no,c.name contract_name
                FROM investment_items i JOIN investment_plans p ON p.id=i.plan_id
                LEFT JOIN projects pr ON pr.id=i.project_id LEFT JOIN contracts c ON c.id=i.contract_id
                WHERE {' AND '.join(where)} ORDER BY p.plan_year DESC,p.id DESC,i.id""", params
        ).fetchall()
        data = [_item_dict(row) for row in rows]
    return {"code": 0, "data": data}


@router.put("/items/{item_id}")
def update_item(item_id: int, payload: ItemPayload):
    start, end = _iso_date(payload.start_date, "开始日期"), _iso_date(payload.end_date, "结束日期")
    if start and end and start > end: raise BusinessError(400, "INV-4001", "投入项开始日期不能晚于结束日期")
    with connect() as conn:
        row = conn.execute("SELECT i.*,p.status plan_status FROM investment_items i JOIN investment_plans p ON p.id=i.plan_id WHERE i.id=?", (item_id,)).fetchone()
        if not row: raise BusinessError(404, "INV-4040", "投入明细不存在")
        if row["plan_status"] not in ("草稿", "已驳回"): raise BusinessError(409, "INV-4091", "已生效投入需通过调整申请变更")
        category_name, subcategory_name = payload.category_name, payload.subcategory_name
        if payload.category_id:
            category = conn.execute("SELECT * FROM investment_categories WHERE id=? AND status='启用'", (payload.category_id,)).fetchone()
            if not category: raise BusinessError(400, "INV-4001", "投入分类不存在或已停用")
            category_name, subcategory_name = category["category_name"], category["subcategory_name"]
        approved = payload.application_amount if payload.approved_amount is None else payload.approved_amount
        conn.execute(
            """UPDATE investment_items SET item_name=?,is_new=?,category_id=?,category_name=?,subcategory_name=?,custom_tags_json=?,
               quantity=?,unit=?,application_amount=?,approved_amount=?,payer=?,business_purpose=?,is_unplanned_reserve=?,
               project_id=?,contract_id=?,start_date=?,end_date=?,planned_payment_amount=?,updated_at=? WHERE id=?""",
            (payload.item_name, int(payload.is_new), payload.category_id, category_name, subcategory_name,
             json.dumps(payload.custom_tags, ensure_ascii=False), payload.quantity, payload.unit, payload.application_amount, approved,
             payload.payer, payload.business_purpose, int(payload.is_unplanned_reserve), payload.project_id, payload.contract_id,
             start, end, payload.planned_payment_amount, now_iso(), item_id),
        )
        _sync_plan_totals(conn, row["plan_id"])
    return {"code": 0, "message": "投入明细已更新"}


@router.delete("/items/{item_id}")
def delete_item(item_id: int):
    with connect() as conn:
        row = conn.execute("SELECT i.plan_id,p.status FROM investment_items i JOIN investment_plans p ON p.id=i.plan_id WHERE i.id=?", (item_id,)).fetchone()
        if not row: raise BusinessError(404, "INV-4040", "投入明细不存在")
        if row["status"] not in ("草稿", "已驳回"): raise BusinessError(409, "INV-4091", "只能删除草稿计划的明细")
        conn.execute("DELETE FROM investment_items WHERE id=?", (item_id,)); _sync_plan_totals(conn, row["plan_id"])
    return {"code": 0, "message": "投入明细已删除"}


@router.post("/plans/{plan_id}/submit")
def submit_plan(plan_id: int):
    with connect() as conn:
        plan = conn.execute("SELECT * FROM investment_plans WHERE id=?", (plan_id,)).fetchone()
        if not plan: raise BusinessError(404, "INV-4040", "投入计划不存在")
        if plan["status"] not in ("草稿", "已驳回"): raise BusinessError(409, "INV-4091", "当前状态不能提交")
        count = conn.execute("SELECT COUNT(*) c FROM investment_items WHERE plan_id=?", (plan_id,)).fetchone()["c"]
        if not count: raise BusinessError(422, "INV-4220", "至少需要一条投入明细")
        ts = now_iso(); conn.execute("UPDATE investment_plans SET status='审批中',current_node=?,submitted_at=?,updated_at=? WHERE id=?", (PLAN_NODES[0], ts, ts, plan_id))
    return {"code": 0, "message": "投入计划已提交审批"}


def _approve_plan(conn, plan_id: int, payload: ApprovalPayload, request: Request):
    plan = conn.execute("SELECT * FROM investment_plans WHERE id=?", (plan_id,)).fetchone()
    if not plan: raise BusinessError(404, "INV-4040", "投入计划不存在")
    if plan["status"] != "审批中" or plan["current_node"] not in PLAN_NODES: raise BusinessError(409, "INV-4091", "投入计划不在可审批节点")
    if payload.action not in ("通过", "驳回"): raise BusinessError(400, "INV-4001", "审批动作仅支持通过或驳回")
    _require_role(request, *PLAN_ROLES[plan["current_node"]])
    actor, role = _actor(request); ts = now_iso(); node = plan["current_node"]
    conn.execute("INSERT INTO investment_approvals(plan_id,node,role,approver,action,comment,created_at) VALUES(?,?,?,?,?,?,?)", (plan_id, node, role, actor, payload.action, payload.comment, ts))
    if payload.action == "驳回":
        status, next_node = "已驳回", "申请人修改"
    else:
        index = PLAN_NODES.index(node)
        if index + 1 < len(PLAN_NODES): status, next_node = "审批中", PLAN_NODES[index + 1]
        else: status, next_node = "待财务确认", "财务确认"
    conn.execute("UPDATE investment_plans SET status=?,current_node=?,approved_at=?,updated_at=? WHERE id=?", (status, next_node, ts if status == "待财务确认" else None, ts, plan_id))
    return {"id": plan_id, "status": status, "current_node": next_node}


@router.get("/approvals/pending")
def pending_approvals(request: Request):
    roles = set((getattr(request.state, "auth_user", None) or {}).get("role_codes") or [request.headers.get("X-Role", "")])
    if "admin" in roles: nodes = list(PLAN_NODES)
    else: nodes = [node for node, allowed in PLAN_ROLES.items() if roles.intersection(allowed)]
    if not nodes: return {"code": 0, "data": []}
    placeholders = ",".join("?" for _ in nodes)
    with connect() as conn:
        data = [dict(r) for r in conn.execute(f"SELECT * FROM investment_plans WHERE status='审批中' AND current_node IN ({placeholders}) ORDER BY submitted_at", nodes)]
    return {"code": 0, "data": data}


@router.post("/plans/{plan_id}/approve")
def approve_plan(plan_id: int, payload: ApprovalPayload, request: Request):
    with connect() as conn: data = _approve_plan(conn, plan_id, payload, request)
    return {"code": 0, "message": f"投入计划已{payload.action}", "data": data}


@router.post("/approvals/batch")
def batch_approve(payload: BatchApprovalPayload, request: Request):
    succeeded, failed = [], []
    with connect() as conn:
        for plan_id in payload.ids:
            try: _approve_plan(conn, plan_id, payload, request); succeeded.append(plan_id)
            except BusinessError as exc: failed.append({"id": plan_id, "message": exc.message})
    return {"code": 0, "message": f"批量处理完成：成功{len(succeeded)}条，失败{len(failed)}条", "data": {"succeeded": succeeded, "failed": failed}}


@router.post("/finance/confirm-batch")
def finance_confirm_batch(payload: BatchApprovalPayload, request: Request):
    _require_role(request, "finance", "admin")
    if payload.action not in ("通过", "驳回"): raise BusinessError(400, "INV-4001", "财务确认仅支持通过或驳回")
    succeeded, failed, ts = [], [], now_iso()
    with connect() as conn:
        actor, role = _actor(request)
        for plan_id in payload.ids:
            plan = conn.execute("SELECT * FROM investment_plans WHERE id=?", (plan_id,)).fetchone()
            if not plan or plan["status"] != "待财务确认": failed.append({"id": plan_id, "message": "非待财务确认状态"}); continue
            status = "已生效" if payload.action == "通过" else "已驳回"
            node = "已生效" if payload.action == "通过" else "申请人修改"
            conn.execute("UPDATE investment_plans SET status=?,current_node=?,finance_confirmed_at=?,updated_at=? WHERE id=?", (status, node, ts if status == "已生效" else None, ts, plan_id))
            if status == "已生效": conn.execute("UPDATE investment_items SET status='执行中',updated_at=? WHERE plan_id=?", (ts, plan_id))
            conn.execute("INSERT INTO investment_approvals(plan_id,node,role,approver,action,comment,created_at) VALUES(?,?,?,?,?,?,?)", (plan_id, "财务确认", role, actor, payload.action, payload.comment, ts))
            succeeded.append(plan_id)
        refresh_investment_warnings(conn)
    return {"code": 0, "message": f"财务确认完成{len(succeeded)}条", "data": {"succeeded": succeeded, "failed": failed}}


@router.get("/finance/export")
def export_finance(status: str = "待财务确认"):
    with connect() as conn:
        rows = conn.execute(
            """SELECT p.plan_no,p.plan_name,p.plan_year,p.department,p.applicant,p.status,i.item_no,i.item_name,
               i.category_name,i.subcategory_name,i.quantity,i.unit,i.application_amount,i.approved_amount,i.payer,i.business_purpose
               FROM investment_plans p JOIN investment_items i ON i.plan_id=p.id WHERE p.status=? ORDER BY p.id,i.id""", (status,)
        ).fetchall()
    wb = Workbook(); ws = wb.active; ws.title = "数字化投入财务复核"
    headers = ["计划编号", "计划名称", "年度", "部门", "申请人", "状态", "明细编号", "投入项", "类别", "子类别", "数量", "单位", "申请金额", "审核金额", "支付方", "业务用途"]
    ws.append(headers)
    for row in rows: ws.append([row[key] for key in row])
    for cell in ws[1]:
        font = copy(cell.font); font.bold = True; cell.font = font
    output = io.BytesIO(); wb.save(output); output.seek(0)
    filename = f"digital-investment-finance-{datetime.now():%Y%m%d}.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename={filename}"})


@router.get("/adjustments")
def list_adjustments():
    with connect() as conn:
        data = [dict(r) for r in conn.execute(
            """SELECT a.*,p.plan_no,p.plan_name,i.item_no,i.item_name FROM investment_adjustments a
               JOIN investment_plans p ON p.id=a.plan_id JOIN investment_items i ON i.id=a.item_id ORDER BY a.id DESC"""
        )]
    return {"code": 0, "data": data}


@router.post("/adjustments")
def create_adjustment(payload: AdjustmentPayload, request: Request):
    actor, _ = _actor(request)
    with connect() as conn:
        item = conn.execute("""SELECT i.*,p.status plan_status,p.plan_year FROM investment_items i JOIN investment_plans p ON p.id=i.plan_id WHERE i.id=? AND i.plan_id=?""", (payload.item_id, payload.plan_id)).fetchone()
        if not item: raise BusinessError(404, "INV-4040", "投入项不存在")
        if item["plan_status"] != "已生效": raise BusinessError(409, "INV-4091", "只有已生效投入可发起调整")
        if payload.requested_amount < float(item["written_off_amount"] or 0): raise BusinessError(422, "INV-4220", "调整后金额不能小于已核销金额")
        ts = now_iso(); no = _next_no(conn, "investment_adjustments", "adjustment_no", f"ADJ-{item['plan_year']}-")
        cur = conn.execute(
            """INSERT INTO investment_adjustments(adjustment_no,plan_id,item_id,adjustment_type,original_amount,requested_amount,
               amount_delta,scope_before,scope_after,reason,applicant,status,current_node,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,'草稿','草稿',?,?)""",
            (no, payload.plan_id, payload.item_id, payload.adjustment_type, item["approved_amount"], payload.requested_amount,
             payload.requested_amount - float(item["approved_amount"] or 0), item["business_purpose"], payload.scope_after,
             payload.reason, actor, ts, ts),
        )
    return {"code": 0, "message": "调整申请已创建", "data": {"id": cur.lastrowid, "adjustment_no": no}}


@router.post("/adjustments/{adjustment_id}/submit")
def submit_adjustment(adjustment_id: int):
    with connect() as conn:
        row = conn.execute("SELECT status FROM investment_adjustments WHERE id=?", (adjustment_id,)).fetchone()
        if not row: raise BusinessError(404, "INV-4040", "调整申请不存在")
        if row["status"] not in ("草稿", "已驳回"): raise BusinessError(409, "INV-4091", "当前状态不能提交")
        conn.execute("UPDATE investment_adjustments SET status='审批中',current_node='部门负责人审批',submitted_at=?,updated_at=? WHERE id=?", (now_iso(), now_iso(), adjustment_id))
    return {"code": 0, "message": "调整申请已提交"}


@router.post("/adjustments/{adjustment_id}/approve")
def approve_adjustment(adjustment_id: int, payload: ApprovalPayload, request: Request):
    with connect() as conn:
        row = conn.execute("SELECT * FROM investment_adjustments WHERE id=?", (adjustment_id,)).fetchone()
        if not row: raise BusinessError(404, "INV-4040", "调整申请不存在")
        if row["status"] != "审批中" or row["current_node"] not in ADJUST_ROLES: raise BusinessError(409, "INV-4091", "调整申请不在可审批节点")
        _require_role(request, *ADJUST_ROLES[row["current_node"]])
        if payload.action not in ("通过", "驳回"): raise BusinessError(400, "INV-4001", "审批动作仅支持通过或驳回")
        actor, role = _actor(request); ts = now_iso(); node = row["current_node"]
        conn.execute("INSERT INTO investment_adjustment_approvals(adjustment_id,node,role,approver,action,comment,created_at) VALUES(?,?,?,?,?,?,?)", (adjustment_id, node, role, actor, payload.action, payload.comment, ts))
        if payload.action == "驳回": status, next_node = "已驳回", "申请人修改"
        elif node == "部门负责人审批": status, next_node = "审批中", "财务审批"
        elif node == "财务审批" and abs(float(row["amount_delta"])) > 50000: status, next_node = "审批中", "分管领导审批"
        else: status, next_node = "已生效", "已生效"
        conn.execute("UPDATE investment_adjustments SET status=?,current_node=?,approved_at=?,updated_at=? WHERE id=?", (status, next_node, ts if status == "已生效" else None, ts, adjustment_id))
        if status == "已生效":
            conn.execute("""UPDATE investment_items SET approved_amount=?,business_purpose=CASE WHEN ?='' THEN business_purpose ELSE ? END,
                           baseline_version=baseline_version+1,updated_at=? WHERE id=?""", (row["requested_amount"], row["scope_after"], row["scope_after"], ts, row["item_id"]))
            _sync_plan_totals(conn, row["plan_id"]); refresh_investment_warnings(conn)
    return {"code": 0, "message": f"调整申请已{payload.action}", "data": {"id": adjustment_id, "status": status, "current_node": next_node}}


@router.get("/execution")
def list_execution(year: Optional[int] = None, keyword: str = ""):
    where, params = ["p.status='已生效'"], []
    if year: where.append("p.plan_year=?"); params.append(year)
    if keyword: where.append("(i.item_no LIKE ? OR i.item_name LIKE ? OR p.plan_no LIKE ?)"); params.extend([f"%{keyword}%"] * 3)
    with connect() as conn:
        refresh_investment_warnings(conn)
        rows = conn.execute(
            f"""SELECT i.*,p.plan_no,p.plan_name,p.plan_year,p.department,p.project_no linked_plan_project,
                pr.project_no,pr.name project_name,c.contract_no,c.name contract_name
                FROM investment_items i JOIN investment_plans p ON p.id=i.plan_id
                LEFT JOIN projects pr ON pr.id=i.project_id LEFT JOIN contracts c ON c.id=i.contract_id
                WHERE {' AND '.join(where)} ORDER BY p.plan_year DESC,i.id DESC""".replace(",p.project_no linked_plan_project", ""), params
        ).fetchall()
        data = [_item_dict(row) for row in rows]
        projects = [dict(r) for r in conn.execute("SELECT id,project_no,name FROM projects ORDER BY id DESC")]
        contracts = [dict(r) for r in conn.execute("SELECT id,contract_no,name,project_id,total_amount,status FROM contracts ORDER BY id DESC")]
    return {"code": 0, "data": data, "meta": {"projects": projects, "contracts": contracts}}


@router.put("/items/{item_id}/binding")
def update_binding(item_id: int, payload: BindingPayload, request: Request):
    _require_role(request, "project_manager", "finance", "product_manager", "admin")
    with connect() as conn:
        row = conn.execute("SELECT * FROM investment_items WHERE id=?", (item_id,)).fetchone()
        if not row: raise BusinessError(404, "INV-4040", "投入项不存在")
        if payload.project_id and not conn.execute("SELECT 1 FROM projects WHERE id=?", (payload.project_id,)).fetchone(): raise BusinessError(400, "INV-4001", "关联项目不存在")
        if payload.contract_id:
            contract = conn.execute("SELECT * FROM contracts WHERE id=?", (payload.contract_id,)).fetchone()
            if not contract: raise BusinessError(400, "INV-4001", "关联合同不存在")
            if payload.project_id and contract["project_id"] and contract["project_id"] != payload.project_id: raise BusinessError(422, "INV-4220", "所选合同与关联项目不一致")
        planned = row["planned_payment_amount"] if payload.planned_payment_amount is None else payload.planned_payment_amount
        conn.execute("UPDATE investment_items SET project_id=?,contract_id=?,planned_payment_amount=?,updated_at=? WHERE id=?", (payload.project_id, payload.contract_id, planned, now_iso(), item_id))
        refresh_investment_warnings(conn)
    return {"code": 0, "message": "项目成本、合同与付款计划关联已更新"}


@router.get("/payments")
def list_payments(item_id: Optional[int] = None):
    with connect() as conn:
        where, params = ("WHERE p.item_id=?", [item_id]) if item_id else ("", [])
        data = [dict(r) for r in conn.execute(f"""SELECT p.*,i.item_no,i.item_name,c.contract_no FROM investment_payments p
             JOIN investment_items i ON i.id=p.item_id LEFT JOIN contracts c ON c.id=p.contract_id {where} ORDER BY p.id DESC""", params)]
    return {"code": 0, "data": data}


@router.post("/payments")
def create_payment(payload: PaymentPayload, request: Request):
    _require_role(request, "finance", "project_manager", "admin")
    payment_date = _iso_date(payload.payment_date, "付款日期")
    if int(payment_date[:4]) != payload.payment_year: raise BusinessError(422, "INV-4220", "付款年份必须与付款日期一致")
    with connect() as conn:
        item = conn.execute("""SELECT i.*,p.plan_year,p.status plan_status FROM investment_items i JOIN investment_plans p ON p.id=i.plan_id WHERE i.id=?""", (payload.item_id,)).fetchone()
        if not item: raise BusinessError(404, "INV-4040", "投入项不存在")
        if item["plan_status"] != "已生效": raise BusinessError(409, "INV-4091", "只有已生效投入可办理核销")
        contract_id = payload.contract_id or item["contract_id"]
        if payload.payment_type == "合同付款":
            if not contract_id: raise BusinessError(422, "INV-4220", "合同付款必须关联合同")
            contract = conn.execute("SELECT * FROM contracts WHERE id=?", (contract_id,)).fetchone()
            if not contract: raise BusinessError(422, "INV-4220", "关联合同不存在")
            if payload.payment_year < item["plan_year"]: raise BusinessError(422, "INV-4220", "合同付款年份不能早于投入年份")
        elif payload.payment_year != item["plan_year"]:
            raise BusinessError(422, "INV-4220", "普通费用付款必须匹配当年投入项")
        remaining = float(item["approved_amount"] or 0) - float(item["written_off_amount"] or 0)
        if payload.amount > remaining + 0.005: raise BusinessError(422, "INV-4220", f"核销金额超过剩余可用金额 {remaining:.2f} 元")
        if conn.execute("SELECT 1 FROM investment_payments WHERE document_no=?", (payload.document_no,)).fetchone(): raise BusinessError(409, "INV-4090", "该付款单据已核销，请勿重复登记")
        ts = now_iso(); no = _next_no(conn, "investment_payments", "payment_no", f"PAY-{payload.payment_year}-")
        cur = conn.execute(
            """INSERT INTO investment_payments(payment_no,item_id,contract_id,payment_type,investment_year,payment_year,amount,
               payment_date,document_no,payer,status,writeoff_amount,description,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,'已核销',?,?,?,?)""",
            (no, payload.item_id, contract_id, payload.payment_type, item["plan_year"], payload.payment_year, payload.amount,
             payment_date, payload.document_no, payload.payer, payload.amount, payload.description, ts, ts),
        )
        written = float(item["written_off_amount"] or 0) + payload.amount
        progress = min(100, written / max(float(item["approved_amount"] or 0), 1) * 100)
        status = "已完成" if progress >= 99.999 else "执行中"
        conn.execute("UPDATE investment_items SET paid_amount=paid_amount+?,written_off_amount=?,progress=?,status=?,updated_at=? WHERE id=?", (payload.amount, written, progress, status, ts, payload.item_id))
        refresh_investment_warnings(conn)
    return {"code": 0, "message": "付款单据已登记并完成投入核销", "data": {"id": cur.lastrowid, "payment_no": no}}


@router.get("/warnings")
def list_warnings(status: str = "待处理"):
    with connect() as conn:
        refresh_investment_warnings(conn)
        where, params = ("WHERE w.status=?", [status]) if status else ("", [])
        data = [dict(r) for r in conn.execute(f"""SELECT w.*,p.plan_no,p.plan_name,i.item_no,i.item_name FROM investment_warnings w
             LEFT JOIN investment_plans p ON p.id=w.plan_id LEFT JOIN investment_items i ON i.id=w.item_id {where} ORDER BY w.triggered_at DESC""", params)]
        rules = [dict(r) for r in conn.execute("SELECT * FROM investment_warning_rules ORDER BY code")]
        for rule in rules: rule["enabled"] = bool(rule["enabled"])
    return {"code": 0, "data": data, "rules": rules}


@router.put("/warning-rules/{code}")
def update_warning_rule(code: str, payload: WarningRulePayload, request: Request):
    _require_role(request, "finance", "admin")
    with connect() as conn:
        if not conn.execute("SELECT 1 FROM investment_warning_rules WHERE code=?", (code,)).fetchone(): raise BusinessError(404, "INV-4040", "预警规则不存在")
        conn.execute("UPDATE investment_warning_rules SET threshold_value=?,days_value=?,enabled=?,level=?,updated_at=? WHERE code=?", (payload.threshold_value, payload.days_value, int(payload.enabled), payload.level, now_iso(), code))
        refresh_investment_warnings(conn)
    return {"code": 0, "message": "预警规则已更新"}


@router.post("/warnings/{warning_id}/resolve")
def resolve_warning(warning_id: int):
    with connect() as conn:
        if not conn.execute("SELECT 1 FROM investment_warnings WHERE id=?", (warning_id,)).fetchone(): raise BusinessError(404, "INV-4040", "预警不存在")
        conn.execute("UPDATE investment_warnings SET status='已处理',resolved_at=? WHERE id=?", (now_iso(), warning_id))
    return {"code": 0, "message": "预警已处理"}


@router.get("/analytics")
def analytics(year: Optional[int] = None):
    with connect() as conn:
        refresh_investment_warnings(conn)
        where, params = ("WHERE p.plan_year=?", [year]) if year else ("", [])
        rows = conn.execute(f"""SELECT p.plan_year,p.department,p.status,i.category_name,i.subcategory_name,i.is_new,
            i.is_unplanned_reserve,i.application_amount,i.approved_amount,i.written_off_amount,i.progress
            FROM investment_plans p LEFT JOIN investment_items i ON i.plan_id=p.id {where}""", params).fetchall()
        plan_count = conn.execute(f"SELECT COUNT(*) c FROM investment_plans p {where}", params).fetchone()["c"]
        warning_count = conn.execute("SELECT COUNT(*) c FROM investment_warnings WHERE status='待处理'").fetchone()["c"]
    totals = {"plan_count": plan_count, "item_count": 0, "application_total": 0, "approved_total": 0, "written_off_total": 0, "remaining_total": 0, "warning_count": warning_count}
    groups = {"year": {}, "category": {}, "department": {}, "status": {}}
    for row in rows:
        if row["category_name"] is None: continue
        totals["item_count"] += 1
        for key, field in (("application_total", "application_amount"), ("approved_total", "approved_amount"), ("written_off_total", "written_off_amount")): totals[key] += float(row[field] or 0)
        for group, field in (("year", "plan_year"), ("category", "category_name"), ("department", "department"), ("status", "status")):
            label = str(row[field] or "未分类"); bucket = groups[group].setdefault(label, {"count": 0, "application": 0, "approved": 0, "written_off": 0})
            bucket["count"] += 1; bucket["application"] += float(row["application_amount"] or 0); bucket["approved"] += float(row["approved_amount"] or 0); bucket["written_off"] += float(row["written_off_amount"] or 0)
    totals["remaining_total"] = max(0, totals["approved_total"] - totals["written_off_total"])
    totals["execution_rate"] = round(totals["written_off_total"] / totals["approved_total"] * 100, 2) if totals["approved_total"] else 0
    return {"code": 0, "data": {"totals": totals, "groups": groups}}
